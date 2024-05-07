from pathlib import Path

import openpyxl as xl
import pytest
import yaml
from fastapi import status
from fastapi.testclient import TestClient

from api import db
from api.main import app
from api.schemas.meta import Meta
from api.schemas.person import Teacher, TeacherBase
from api.schemas.timeslot import (
    MonthlyAttendance,
    TimeslotJS,
    UpdateAttendanceReq,
)

client = TestClient(app)


def read_timetable_excel(file_path: str, month: int):
    wb = xl.load_workbook(file_path)
    ret = []
    for sheet in wb.sheetnames:
        if f"{month}月" not in sheet:
            continue
        ws = wb[sheet]
        for i in range(100):

            block_row = ws[i * 3 + 1 : (i + 1) * 3 + 1][:100]
            block_row_values = []
            for row in block_row:
                block_row_values.append([cell.value for cell in row])
            ret.append(block_row_values)
    return ret


@pytest.fixture(autouse=True)
def set_db() -> None:
    setattr(db.DBModelBase.Meta, "host", "http://localhost:8000")
    db.DBModelBase.create_table(
        read_capacity_units=10, write_capacity_units=10, wait=True
    )
    yield
    db.DBModelBase.delete_table()


def test_meta():
    res_1 = client.post("/metas/", json={"school_name": "テスト校"})
    assert res_1.status_code == status.HTTP_201_CREATED
    meta_1 = Meta(**res_1.json())

    assert meta_1 == Meta(**client.get(f"/metas/{meta_1.school_id}").json())
    assert client.get("/metas/hogehoge").status_code == status.HTTP_404_NOT_FOUND
    assert (
        client.post("/metas/", json={"school_name": "テスト校"}).status_code
        == status.HTTP_409_CONFLICT
    )

    res_2 = client.post("/metas/", json={"school_name": "テスト校2"})
    assert res_2.status_code == status.HTTP_201_CREATED
    meta_2 = Meta(**res_2.json())

    assert [meta_1, meta_2] == [
        Meta(**meta_json) for meta_json in client.get("/metas/").json()
    ]

    res_3 = client.put(f"/metas/{meta_1.school_id}", json={"school_name": "テスト校3"})
    assert res_3.status_code == status.HTTP_200_OK
    meta_3 = Meta(**res_3.json())
    assert meta_3 == Meta(**client.get(f"/metas/{meta_1.school_id}").json())

    assert (
        client.put("/metas/hogehoge", json={"school_name": "テスト校4"}).status_code
        == status.HTTP_404_NOT_FOUND
    )

    assert client.delete(f"/metas/{meta_2.school_id}").status_code == status.HTTP_200_OK
    assert [meta_3] == [Meta(**meta_json) for meta_json in client.get("/metas/").json()]

    assert client.delete("/metas/hogehoge").status_code == status.HTTP_404_NOT_FOUND


def test_teacher():
    school_id = Meta(
        **client.post("/metas/", json={"school_name": "テスト校"}).json()
    ).school_id

    teacher_base_1 = TeacherBase(
        display_name="test",
        given_name="test",
        family_name="test",
        school_id=school_id,
        teacher_type="teacher",
        lecture_hourly_pay=1000,
        office_hourly_pay=1000,
        sub="sub",
    )
    res_1 = client.post("/teachers/", json=teacher_base_1.model_dump())
    assert res_1.status_code == status.HTTP_201_CREATED
    teacher_1 = Teacher(**res_1.json())

    assert teacher_1 == Teacher(**client.get(f"/teachers/{teacher_1.id}").json())
    assert teacher_1 == Teacher(**client.get(f"/teachers/sub/{teacher_1.sub}").json())
    assert (
        Teacher(**client.get("/teachers/sub/hogehoge").json()).display_name == "Guest"
    )
    assert client.get("/teachers/hogehoge").status_code == status.HTTP_404_NOT_FOUND

    assert (
        client.post("/teachers/", json=teacher_base_1.model_dump()).status_code
        == status.HTTP_409_CONFLICT
    )

    teacher_base_2 = TeacherBase(
        display_name="test2",
        given_name="test2",
        family_name="test2",
        school_id=school_id,
        teacher_type="teacher",
        lecture_hourly_pay=1000,
        office_hourly_pay=1000,
    )
    res_2 = client.post("/teachers/", json=teacher_base_2.model_dump())
    assert res_2.status_code == status.HTTP_201_CREATED
    teacher_2 = Teacher(**res_2.json())

    assert sorted([teacher_1, teacher_2]) == sorted(
        [
            Teacher(**teacher_json)
            for teacher_json in client.get(f"/teachers/bulk/{school_id}").json()
        ],
    )

    teacher_base_3 = TeacherBase(
        display_name="test3",
        given_name="test3",
        family_name="test3",
        school_id=school_id,
        teacher_type="teacher",
        lecture_hourly_pay=1000,
        office_hourly_pay=1000,
    )
    res_3 = client.put(f"/teachers/{teacher_1.id}", json=teacher_base_3.model_dump())
    assert res_3.status_code == status.HTTP_200_OK
    teacher_3 = Teacher(**res_3.json())
    assert teacher_3 == Teacher(**client.get(f"/teachers/{teacher_1.id}").json())

    assert (
        client.put("/teachers/hogehoge", json=teacher_base_3.model_dump()).status_code
        == status.HTTP_404_NOT_FOUND
    )

    assert client.delete(f"/teachers/{teacher_2.id}").status_code == status.HTTP_200_OK
    assert [teacher_3] == [
        Teacher(**teacher_json)
        for teacher_json in client.get(f"/teachers/bulk/{school_id}").json()
    ]

    assert client.delete("/teachers/hogehoge").status_code == status.HTTP_404_NOT_FOUND


def test_calc_salary(snapshot):
    snapshot_dir = Path("tests/snapshots")
    snapshot.snapshot_dir = str(snapshot_dir)

    meta = Meta(**client.post("/metas/", json={"school_name": "テスト校"}).json())
    school_id = meta.school_id

    # send csv file
    with open("tests/data/講師情報.csv", "rb") as f:
        res = client.post(
            f"teachers/bulk/{school_id}", files={"file": ("test.csv", f, "text/csv")}
        )
        assert res.status_code == status.HTTP_201_CREATED, res.json()["detail"]

    # get teacher list
    res = client.get(f"teachers/bulk/{school_id}")
    teacher_list = [Teacher(**teacher) for teacher in res.json()]
    snapshot.assert_match(yaml.dump(teacher_list), "teacher_list_from_csv.yml")

    month_1 = 11
    year = 2022
    content = read_timetable_excel("tests/data/2022　新時間割表 10月分.xlsx", month_1)
    meetings = []
    res = client.post(
        f"salary/bulk/{school_id}?year={year}&month={month_1}",
        json={"content": content, "meetings": meetings},
    )
    assert res.status_code == status.HTTP_200_OK, res.json()["detail"]

    month_2 = 10
    client.post(
        f"salary/bulk/{school_id}?year={year}&month={month_2}",
        json={"content": content, "meetings": meetings},
    )

    res = client.get(f"salary/bulk/{school_id}?year={year}&month={month_1}")
    assert res.status_code == status.HTTP_200_OK, res.json()["detail"]
    monthly_attendance_list_between = [
        MonthlyAttendance(**monthly_attendance) for monthly_attendance in res.json()
    ]
    snapshot.assert_match(
        yaml.dump(monthly_attendance_list_between),
        "monthly_attendance_list_from_excel.yml",
    )

    monthly_attendance_0 = monthly_attendance_list_between[0]
    res = client.get(
        f"salary/{monthly_attendance_0.teacher.id}?year={year}&month={month_1}"
    )
    assert res.status_code == status.HTTP_200_OK, res.json()["detail"]
    assert monthly_attendance_0 == MonthlyAttendance(**res.json())


    res = client.get(f"salary/bulk/{school_id}/between?start_year={year}&start_month={month_2}&end_year={year}&end_month={month_1}")
    assert res.status_code == status.HTTP_200_OK, res.json()["detail"]
    monthly_attendance_list_between = [
        MonthlyAttendance(**monthly_attendance) for monthly_attendance in res.json()
    ]
    snapshot.assert_match(
        yaml.dump(monthly_attendance_list_between),
        "monthly_attendance_list_between_from_excel.yml",
    )

    timeslot_js_list = [
        TimeslotJS.from_timeslot(timeslot)
        for timeslot in monthly_attendance_0.timeslot_list
    ]

    req = UpdateAttendanceReq(
        timeslot_js_list=timeslot_js_list,
        extra_payment=1000,
        remark="test",
        teacher=monthly_attendance_0.teacher,
    )

    res = client.put(
        f"salary/{monthly_attendance_0.teacher.id}?year={year}&month={month_1}",
        json=req.model_dump(),
    )
    assert res.status_code == status.HTTP_200_OK, res.json()["detail"]
    updated_monthly_attendance = MonthlyAttendance(**res.json())
    assert updated_monthly_attendance == MonthlyAttendance(
        **client.get(
            f"salary/{monthly_attendance_0.teacher.id}?year={year}&month={month_1}"
        ).json()
    )
    assert updated_monthly_attendance.extra_payment == 1000
    assert updated_monthly_attendance.remark == "test"

    res = client.delete(f"salary/bulk/{school_id}?year={year}&month={month_1}")
    assert res.status_code == status.HTTP_200_OK, res.json()["detail"]
    assert client.get(f"salary/bulk/{school_id}?year={year}&month={month_1}").json() == []
    
